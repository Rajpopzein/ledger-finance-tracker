import csv
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
import xlrd
from pypdf import PdfReader

APP_LABELS = {
    "google_pay": "Google Pay",
    "phonepe": "PhonePe",
}

DATE_KEYS = [
    "date", "transaction date", "txn date", "timestamp", "time",
    "transaction time", "paid on", "created at", "completed at", "last updated time"
]
AMOUNT_KEYS = [
    "amount", "transaction amount", "txn amount", "amount paid",
    "total amount", "value", "total price"
]
DESC_KEYS = [
    "description", "merchant", "merchant name", "name", "payee", "payee name",
    "recipient", "recipient name", "paid to", "received from", "title",
    "counterparty", "transaction note", "note"
]
TYPE_KEYS = [
    "type", "direction", "transaction type", "txn type", "payment type",
    "flow", "debit credit", "dr cr"
]
STATUS_KEYS = ["status", "transaction status", "payment status", "state"]
REF_KEYS = [
    "upi transaction id", "upi txn id", "upi id", "upi rrn", "rrn", "utr",
    "transaction id", "txn id", "reference id", "reference", "ref no", "ref id"
]

SUCCESS_WORDS = {"success", "successful", "completed", "complete", "paid", "received"}
FAIL_WORDS = {"failed", "failure", "declined", "pending", "processing", "cancelled", "canceled", "expired", "reversed"}

def app_label(app: str) -> str:
    if app not in APP_LABELS:
        raise ValueError("Supported UPI apps are Google Pay and PhonePe")
    return APP_LABELS[app]

def _norm(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()

def _key(row, candidates):
    lookup = {_norm(k): k for k in row.keys() if k is not None}
    for candidate in candidates:
        key = _norm(candidate)
        if key in lookup:
            return lookup[key]
    return None

def _decimal(value):
    if value in (None, ""):
        return None
    raw = str(value).strip()
    negative = raw.startswith("(") and raw.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", raw)
    try:
        amount = Decimal(cleaned)
        return -amount if negative else amount
    except (InvalidOperation, ValueError):
        return None

def _date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    raw = str(value or "").strip()
    if not raw:
        raise ValueError("Missing transaction date")

    iso_candidate = raw.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(iso_candidate)
    except ValueError:
        pass

    formats = (
        "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y",
        "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%d %b %y",
        "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S",
        "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M",
        "%d %b %Y, %I:%M %p", "%d %b %Y %I:%M %p",
        "%b %d, %Y, %I:%M:%S %p", "%b %d, %Y %I:%M:%S %p",
        "%b %d, %Y, %I:%M %p", "%b %d, %Y %I:%M %p",
        "%d %B %Y, %I:%M %p", "%d %B %Y %I:%M %p",
        "%B %d, %Y, %I:%M %p", "%B %d, %Y %I:%M %p",
        "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y",
        "%d %b %Y %H:%M", "%d %B %Y %H:%M",
    )
    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    raise ValueError(f"Unsupported transaction date: {raw}")

def _status_ok(value):
    raw = _norm(value)
    if not raw:
        return True
    tokens = set(raw.split())
    if tokens & FAIL_WORDS:
        return False
    if tokens & SUCCESS_WORDS:
        return True
    return True

def _direction(*values):
    raw = " ".join(_norm(v) for v in values if v not in (None, ""))
    if not raw:
        return None

    credit_phrases = (
        "received from", "money received", "received", "credit",
        "cashback", "refund", "refunded", "reward"
    )
    debit_phrases = (
        "paid to", "sent to", "money sent", "debit",
        "payment to", "paid", "sent", "purchase"
    )

    if any(phrase in raw for phrase in credit_phrases):
        return "credit"
    if any(phrase in raw for phrase in debit_phrases):
        return "debit"

    tokens = set(raw.split())
    if "cr" in tokens:
        return "credit"
    if "dr" in tokens:
        return "debit"
    return None

def _reference(value):
    raw = str(value or "").strip()
    if not raw:
        return None

    # Prefer UPI/UTR/RRN-like numeric references but keep non-numeric IDs too.
    match = re.search(r"(?<!\d)(\d{9,18})(?!\d)", raw)
    if match:
        return match.group(1)
    return raw[:80]

def _matrix_rows(matrix):
    if not matrix:
        return []

    aliases = {_norm(x) for x in DATE_KEYS + AMOUNT_KEYS + DESC_KEYS + REF_KEYS}
    header_index = None
    for index, values in enumerate(matrix[:40]):
        normalized = {_norm(v) for v in values if v not in (None, "")}
        date_hit = bool(normalized & {_norm(x) for x in DATE_KEYS})
        amount_hit = bool(normalized & {_norm(x) for x in AMOUNT_KEYS})
        useful_hits = len(normalized & aliases)
        if date_hit and amount_hit and useful_hits >= 2:
            header_index = index
            break

    if header_index is None:
        raise ValueError("Could not identify UPI transaction columns in this file")

    headers = [str(v).strip() if v is not None else "" for v in matrix[header_index]]
    return [
        dict(zip(headers, values))
        for values in matrix[header_index + 1:]
        if any(v not in (None, "") for v in values)
    ]

def _collect_json_rows(value):
    rows = []

    def walk(node):
        if isinstance(node, list):
            for item in node:
                walk(item)
            return
        if not isinstance(node, dict):
            return

        normalized = {_norm(k) for k in node.keys()}
        has_date = bool(normalized & {_norm(x) for x in DATE_KEYS})
        has_amount = bool(normalized & {_norm(x) for x in AMOUNT_KEYS})
        if has_date and has_amount:
            rows.append(node)

        for child in node.values():
            if isinstance(child, (dict, list)):
                walk(child)

    walk(value)
    return rows


PDF_DATE_PATTERNS = (
    re.compile(r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}(?:\s+\d{1,2}:\d{2}(?::\d{2})?\s*(?:AM|PM)?)?\b", re.I),
    re.compile(r"\b\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4}(?:(?:,\s*|\s+)\d{1,2}:\d{2}(?::\d{2})?\s*(?:AM|PM)?)?\b", re.I),
    re.compile(r"\b[A-Za-z]{3,9}\s+\d{1,2},\s*\d{4}(?:(?:,\s*|\s+)\d{1,2}:\d{2}(?::\d{2})?\s*(?:AM|PM)?)?\b", re.I),
)

def _pdf_date(text):
    for pattern in PDF_DATE_PATTERNS:
        for match in pattern.finditer(text):
            try:
                return _date(match.group(0))
            except ValueError:
                continue
    return None

def _pdf_amount(text):
    patterns = (
        re.compile(r"(?:₹|INR|Rs\.?)\s*([+-]?\s*\d[\d,]*(?:\.\d{1,2})?)", re.I),
        re.compile(r"\bamount(?:\s+paid)?\s*[:\-]?\s*(?:₹|INR|Rs\.?)?\s*([+-]?\s*\d[\d,]*(?:\.\d{1,2})?)", re.I),
    )
    for pattern in patterns:
        for match in pattern.finditer(text):
            value = _decimal(match.group(1))
            if value is not None and value != 0:
                return value
    return None

def _pdf_reference(text):
    labeled = re.search(
        r"(?:UPI\s*(?:transaction|txn)?\s*ID|UTR|RRN|reference\s*ID|transaction\s*ID)\s*[:#\-]?\s*([A-Za-z0-9\-]{6,40})",
        text,
        re.I,
    )
    if labeled:
        return _reference(labeled.group(1))

    if re.search(r"\b(?:UPI|UTR|RRN|reference)\b", text, re.I):
        numeric = re.search(r"(?<!\d)(\d{9,18})(?!\d)", text)
        if numeric:
            return numeric.group(1)
    return None

def _pdf_merchant(lines, label):
    prefixes = (
        "paid to", "sent to", "received from", "merchant",
        "payee", "recipient", "payment to"
    )
    for index, line in enumerate(lines):
        normalized = _norm(line)
        for prefix in prefixes:
            if normalized.startswith(prefix):
                remainder = re.sub(
                    rf"^\s*{re.escape(prefix)}\s*[:\-]?\s*",
                    "",
                    line,
                    flags=re.I,
                ).strip()
                if remainder and _norm(remainder) != prefix:
                    return remainder[:160]
                if index + 1 < len(lines):
                    candidate = lines[index + 1].strip()
                    if candidate and not _pdf_date(candidate) and _pdf_amount(candidate) is None:
                        return candidate[:160]
    return label

GPAY_ROW_RE = re.compile(
    r"^(?P<date>\\d{1,2}\\s*[A-Za-z]{3},\\s*\\d{4})\\s+"
    r"(?P<detail>.+?)\\s+₹\\s*(?P<amount>\\d[\\d,]*(?:\\.\\d{1,2})?)$",
    re.I,
)
GPAY_TIME_REF_RE = re.compile(
    r"^(?P<time>\\d{1,2}:\\d{2}\\s*(?:AM|PM))\\s+"
    r"UPI\\s*Transaction\\s*ID\\s*:\\s*(?P<ref>\\d{9,18})$",
    re.I,
)

def _gpay_detail(detail: str):
    patterns = (
        ("internal_transfer", "debit", re.compile(r"^self\\s*transfer\\s*to\\s*(.+)$", re.I)),
        ("income", "credit", re.compile(r"^received\\s*from\\s*(.+)$", re.I)),
        ("expense", "debit", re.compile(r"^paid\\s*to\\s*(.+)$", re.I)),
    )
    for txn_type, direction, pattern in patterns:
        match = pattern.match(detail)
        if match:
            return txn_type, direction, re.sub(r"\\s+", " ", match.group(1)).strip()
    return None, None, None

def _gpay_date(date_text: str, time_text: str):
    clean_date = re.sub(r"^(\\d{1,2})([A-Za-z]{3})", r"\\1 \\2", date_text.strip())
    clean_time = re.sub(r"(?i)(\\d)(AM|PM)$", r"\\1 \\2", re.sub(r"\\s+", " ", time_text).strip())
    return datetime.strptime(f"{clean_date} {clean_time}", "%d %b, %Y %I:%M %p")

def _gpay_header_totals(text: str):
    flat = re.sub(r"\\s+", " ", text)
    sent_match = re.search(r"\\bSent\\s+₹\\s*([\\d,]+(?:\\.\\d{1,2})?)", flat, re.I)
    received_match = re.search(r"\\bReceived\\s+₹\\s*([\\d,]+(?:\\.\\d{1,2})?)", flat, re.I)
    sent = _decimal(sent_match.group(1)) if sent_match else None
    received = _decimal(received_match.group(1)) if received_match else None
    return sent, received

def _google_pay_pdf_rows(reader: PdfReader, label: str):
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text(extraction_mode="layout") or "")
        except TypeError:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")

    text = "\\n".join(pages).strip()
    if not text:
        raise ValueError(
            "This Google Pay PDF does not contain selectable text. "
            "Please export the transaction statement PDF directly from Google Pay."
        )

    rows = []
    for page_text in pages:
        lines = [
            re.sub(r"\\s+", " ", line).strip()
            for line in page_text.splitlines()
            if line.strip()
        ]
        starts = []
        for index, line in enumerate(lines):
            match = GPAY_ROW_RE.match(line)
            if not match:
                continue
            txn_type, direction, merchant = _gpay_detail(match.group("detail"))
            if txn_type:
                starts.append((index, match, txn_type, direction, merchant))

        for pos, (start_index, match, txn_type, direction, merchant) in enumerate(starts):
            end_index = starts[pos + 1][0] if pos + 1 < len(starts) else min(len(lines), start_index + 7)
            block_lines = lines[start_index:end_index]

            time_ref = None
            for line in block_lines[1:4]:
                candidate = GPAY_TIME_REF_RE.match(line)
                if candidate:
                    time_ref = candidate
                    break
            if not time_ref:
                continue

            try:
                txn_at = _gpay_date(match.group("date"), time_ref.group("time"))
            except ValueError:
                continue

            amount = _decimal(match.group("amount"))
            if amount is None or amount == 0:
                continue

            source_account = None
            destination_account = merchant if txn_type == "internal_transfer" else None
            for line in block_lines[1:]:
                source_match = re.match(r"^paid\\s*by\\s*(.+)$", line, re.I)
                if source_match:
                    source_account = re.sub(r"\\s+", " ", source_match.group(1)).strip()
                    continue
                destination_match = re.match(r"^paid\\s*to\\s*(.+)$", line, re.I)
                if direction == "credit" and destination_match:
                    destination_account = re.sub(r"\\s+", " ", destination_match.group(1)).strip()

            account_hint = source_account if direction == "debit" else destination_account
            reference = time_ref.group("ref")
            description_parts = [
                "Google Pay",
                f"{'received from' if direction == 'credit' else 'paid to'} {merchant}",
            ]
            if account_hint:
                description_parts.append(account_hint)

            rows.append({
                "txn_at": txn_at,
                "amount": abs(amount),
                "direction": direction,
                "txn_type": txn_type,
                "description": " | ".join(description_parts),
                "merchant": merchant[:160],
                "upi_ref": reference,
                "account_hint": account_hint,
                "source_account_hint": source_account,
                "destination_account_hint": destination_account,
                "source_app": "google_pay",
                "source_label": label,
            })

    if not rows:
        raise ValueError(
            "Ledger could read the Google Pay PDF, but no transaction rows matched the supported statement format."
        )

    expected_sent, expected_received = _gpay_header_totals(text)
    parsed_sent = sum(
        (row["amount"] for row in rows if row["direction"] == "debit" and row["txn_type"] != "internal_transfer"),
        Decimal("0"),
    )
    parsed_received = sum(
        (row["amount"] for row in rows if row["direction"] == "credit"),
        Decimal("0"),
    )
    if expected_sent is not None and abs(parsed_sent - expected_sent) > Decimal("0.01"):
        raise ValueError(
            f"Google Pay PDF was only partially parsed: statement Sent total is ₹{expected_sent}, "
            f"but Ledger found ₹{parsed_sent}. No transactions were imported."
        )
    if expected_received is not None and abs(parsed_received - expected_received) > Decimal("0.01"):
        raise ValueError(
            f"Google Pay PDF was only partially parsed: statement Received total is ₹{expected_received}, "
            f"but Ledger found ₹{parsed_received}. No transactions were imported."
        )

    return rows

def _pdf_rows(content: bytes, app: str):
    label = app_label(app)
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ValueError("Could not open this PDF statement") from exc

    page_text = []
    for page in reader.pages:
        try:
            page_text.append(page.extract_text() or "")
        except Exception:
            page_text.append("")

    text = "\n".join(page_text).strip()
    if not text:
        raise ValueError(
            "This PDF does not contain selectable text. "
            "Please export a searchable PDF/e-statement instead of a scanned image PDF."
        )

    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]
    starts = []
    for index, line in enumerate(lines):
        lowered = _norm(line)
        if any(word in lowered for word in ("statement period", "statement from", "from date", "to date")):
            continue
        txn_at = _pdf_date(line)
        if txn_at:
            starts.append((index, txn_at))

    rows = []
    for pos, (start_index, txn_at) in enumerate(starts):
        end_index = starts[pos + 1][0] if pos + 1 < len(starts) else min(len(lines), start_index + 14)
        block_lines = lines[start_index:end_index]
        block = "\n".join(block_lines)

        if not _status_ok(block):
            continue

        amount = _pdf_amount(block)
        direction = _direction(block)
        if amount is None or amount == 0 or direction is None:
            continue

        reference = _pdf_reference(block)
        merchant = _pdf_merchant(block_lines, label)
        rows.append({
            "txn_at": txn_at,
            "amount": abs(amount),
            "direction": direction,
            "description": block[:2000],
            "merchant": merchant,
            "upi_ref": reference,
            "source_app": app,
            "source_label": label,
        })

    if not rows:
        raise ValueError(
            "Ledger could read this PDF, but could not identify supported UPI transaction rows. "
            "Use the app's transaction-history statement PDF, or upload CSV/XLSX/XLS/JSON."
        )

    return rows

def normalize_upi_rows(rows, app: str):
    label = app_label(app)
    out = []
    unresolved = 0

    for row in rows:
        date_key = _key(row, DATE_KEYS)
        amount_key = _key(row, AMOUNT_KEYS)
        desc_key = _key(row, DESC_KEYS)
        type_key = _key(row, TYPE_KEYS)
        status_key = _key(row, STATUS_KEYS)
        ref_key = _key(row, REF_KEYS)

        if not date_key or not amount_key:
            continue
        if status_key and not _status_ok(row.get(status_key)):
            continue

        try:
            txn_at = _date(row.get(date_key))
        except ValueError:
            continue

        raw_amount = _decimal(row.get(amount_key))
        if raw_amount is None or raw_amount == 0:
            continue

        description = str(row.get(desc_key) or "").strip() if desc_key else ""
        type_value = row.get(type_key) if type_key else ""
        direction = _direction(type_value, description)

        if direction is None:
            if raw_amount < 0:
                direction = "debit"
            elif str(row.get(amount_key) or "").strip().startswith("+"):
                direction = "credit"

        if direction is None:
            unresolved += 1
            continue

        reference = _reference(row.get(ref_key)) if ref_key else None
        if not reference:
            # Some exports place the UPI/UTR inside a description field.
            reference = _reference(description) if re.search(r"\b(?:upi|utr|rrn|ref)\b", description, re.I) else None

        merchant = description or label
        out.append({
            "txn_at": txn_at,
            "amount": abs(raw_amount),
            "direction": direction,
            "description": description or f"{label} transaction",
            "merchant": merchant[:160],
            "upi_ref": reference,
            "source_app": app,
            "source_label": label,
        })

    if unresolved and not out:
        raise ValueError(
            "Transactions were found, but Ledger could not determine whether they were sent or received. "
            "Use an export that includes transaction type/direction."
        )

    return out

def parse_upi_statement(app: str, name: str, content: bytes):
    app_label(app)
    lower = name.lower()

    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        matrix = list(csv.reader(io.StringIO(text)))
        return normalize_upi_rows(_matrix_rows(matrix), app)

    if lower.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        matrix = list(ws.iter_rows(values_only=True))
        return normalize_upi_rows(_matrix_rows(matrix), app)

    if lower.endswith(".xls"):
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        matrix = []
        for row_index in range(sheet.nrows):
            values = []
            for cell in sheet.row(row_index):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                else:
                    values.append(cell.value)
            matrix.append(values)
        return normalize_upi_rows(_matrix_rows(matrix), app)

    if lower.endswith(".json"):
        try:
            payload = json.loads(content.decode("utf-8-sig", errors="replace"))
        except json.JSONDecodeError as exc:
            raise ValueError("Invalid JSON export") from exc
        return normalize_upi_rows(_collect_json_rows(payload), app)

    if lower.endswith(".pdf"):
        if app == "google_pay":
            try:
                reader = PdfReader(io.BytesIO(content))
            except Exception as exc:
                raise ValueError("Could not open this Google Pay PDF statement") from exc
            return _google_pay_pdf_rows(reader, app_label(app))
        return _pdf_rows(content, app)

    raise ValueError("UPI app import supports CSV, XLSX, XLS, JSON and PDF files")

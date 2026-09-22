import csv, io, re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

DATE_KEYS = ["date", "transaction date", "txn date", "value date", "value dt", "transaction dt", "txn dt"]
DESC_KEYS = ["description", "narration", "details", "transaction remarks", "remarks", "particulars", "transaction details"]
DEBIT_KEYS = [
    "debit", "withdrawal", "withdrawal amt", "withdrawal amount",
    "withdrawal dr", "debit amount", "debit amt", "dr amount"
]
CREDIT_KEYS = [
    "credit", "deposit", "deposit amt", "deposit amount",
    "deposit cr", "credit amount", "credit amt", "cr amount"
]
AMOUNT_KEYS = ["amount", "transaction amount", "txn amount"]
COMBINED_AMOUNT_KEYS = [
    "debit credit", "withdrawal dr deposit cr", "withdrawal deposit",
    "withdrawal dr deposit cr amount"
]
TYPE_KEYS = ["type", "dr cr", "transaction type", "txn type", "debit credit type"]
REF_KEYS = ["reference", "ref no", "reference no", "utr", "transaction id", "txn id", "chq ref no", "cheque ref no"]

def _norm_header(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()

def _aliases(values):
    return {_norm_header(v) for v in values}

DATE_ALIASES = _aliases(DATE_KEYS)
DESC_ALIASES = _aliases(DESC_KEYS)
DEBIT_ALIASES = _aliases(DEBIT_KEYS)
CREDIT_ALIASES = _aliases(CREDIT_KEYS)
AMOUNT_ALIASES = _aliases(AMOUNT_KEYS)
COMBINED_AMOUNT_ALIASES = _aliases(COMBINED_AMOUNT_KEYS)
TYPE_ALIASES = _aliases(TYPE_KEYS)
REF_ALIASES = _aliases(REF_KEYS)

def _key(row, candidates):
    lookup = {_norm_header(k): k for k in row.keys() if k is not None}
    for candidate in candidates:
        normalized = _norm_header(candidate)
        if normalized in lookup:
            return lookup[normalized]
    return None

def _decimal(value):
    if value in (None, ""):
        return Decimal("0")
    raw = str(value).strip()
    negative = raw.startswith("(") and raw.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", raw)
    try:
        amount = Decimal(cleaned or "0")
        return -amount if negative else amount
    except InvalidOperation:
        return Decimal("0")

def _direction_marker(value):
    raw = str(value or "").strip().lower()
    if not raw:
        return None

    compact = re.sub(r"[^a-z]+", " ", raw).strip()
    tokens = set(compact.split())

    if "dr" in tokens or "debit" in tokens:
        return "debit"
    if "cr" in tokens or "credit" in tokens:
        return "credit"
    if compact in {"d", "db"}:
        return "debit"
    if compact in {"c", "crd"}:
        return "credit"
    return None

def _date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    raw = str(value or "").strip()
    if not raw:
        raise ValueError("Missing transaction date")

    iso_candidate = raw.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(iso_candidate)
    except ValueError:
        pass

    formats = (
        "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y",
        "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%d %b %y",
        "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S",
        "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M",
        "%d-%b-%Y %H:%M:%S", "%d %b %Y %H:%M:%S",
    )
    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass

    raise ValueError(f"Unsupported date: {raw}")

def _looks_like_header(values) -> bool:
    normalized = {_norm_header(v) for v in values if v not in (None, "")}
    has_date = bool(normalized & DATE_ALIASES)
    has_description = bool(normalized & DESC_ALIASES)
    has_amount = bool(
        normalized &
        (DEBIT_ALIASES | CREDIT_ALIASES | AMOUNT_ALIASES | COMBINED_AMOUNT_ALIASES)
    )
    return has_date and has_description and has_amount

def _rows_from_matrix(matrix):
    header_index = None
    for index, values in enumerate(matrix[:30]):
        if _looks_like_header(values):
            header_index = index
            break

    if header_index is None:
        sample = []
        for values in matrix[:10]:
            sample.extend([_norm_header(v) for v in values if v not in (None, "")])
        visible = ", ".join(list(dict.fromkeys(x for x in sample if x))[:12]) if sample else "none"
        raise ValueError(
            "Could not identify the bank statement header row. "
            f"Detected columns/text: {visible}"
        )

    headers = [str(v).strip() if v is not None else "" for v in matrix[header_index]]
    return [
        dict(zip(headers, row))
        for row in matrix[header_index + 1:]
        if any(cell not in (None, "") for cell in row)
    ]

def normalize_rows(rows):
    out = []

    for row in rows:
        date_key = _key(row, DATE_KEYS)
        description_key = _key(row, DESC_KEYS)
        debit_key = _key(row, DEBIT_KEYS)
        credit_key = _key(row, CREDIT_KEYS)
        amount_key = _key(row, AMOUNT_KEYS)
        combined_amount_key = _key(row, COMBINED_AMOUNT_KEYS)
        type_key = _key(row, TYPE_KEYS)
        ref_key = _key(row, REF_KEYS)

        if not date_key or not description_key:
            continue

        amount = Decimal("0")
        direction = None

        if debit_key or credit_key:
            debit_amount = _decimal(row.get(debit_key)) if debit_key else Decimal("0")
            credit_amount = _decimal(row.get(credit_key)) if credit_key else Decimal("0")

            if debit_amount != 0:
                amount, direction = abs(debit_amount), "debit"
            elif credit_amount != 0:
                amount, direction = abs(credit_amount), "credit"

        elif combined_amount_key:
            raw_value = row.get(combined_amount_key)
            parsed_amount = _decimal(raw_value)
            marker = _direction_marker(raw_value)

            if parsed_amount != 0 and marker:
                amount, direction = abs(parsed_amount), marker

        elif amount_key:
            raw_value = row.get(amount_key)
            parsed_amount = _decimal(raw_value)

            if parsed_amount != 0:
                marker = _direction_marker(row.get(type_key)) if type_key else None
                marker = marker or _direction_marker(raw_value)

                if marker:
                    amount, direction = abs(parsed_amount), marker
                else:
                    amount = abs(parsed_amount)
                    direction = "credit" if parsed_amount > 0 and str(raw_value).strip().startswith("+") else "debit"

        if amount == 0 or direction is None:
            continue

        try:
            txn_at = _date(row.get(date_key))
        except ValueError:
            continue

        description = str(row.get(description_key) or "").strip()
        reference = str(row.get(ref_key) or "").strip() if ref_key else None

        out.append({
            "txn_at": txn_at,
            "amount": amount,
            "direction": direction,
            "description": description,
            "bank_ref": reference or None,
        })

    return out

def parse_statement(name: str, content: bytes):
    lower = name.lower()

    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        matrix = list(csv.reader(io.StringIO(text)))
        rows = _rows_from_matrix(matrix)
        return normalize_rows(rows)

    if lower.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        matrix = list(ws.iter_rows(values_only=True))
        if not matrix:
            return []
        rows = _rows_from_matrix(matrix)
        return normalize_rows(rows)

    raise ValueError("Only CSV and XLSX are enabled in v1")
